from pptx import Presentation
from pptx.util import Inches, Pt
from io import BytesIO

def create_presentation(pages, texts, output_path):
    """Genera un PowerPoint con las imágenes de fondo y texto editable."""
    prs = Presentation()

    for i, (page, text) in enumerate(zip(pages, texts), start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Imagen como fondo
        img_buffer = BytesIO()
        page.save(img_buffer, format='PNG')
        img_buffer.seek(0)
        slide.shapes.add_picture(img_buffer, 0, 0,
                                 width=prs.slide_width,
                                 height=prs.slide_height)

        # Texto editable
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5),
                                         prs.slide_width - Inches(1),
                                         prs.slide_height - Inches(1))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(14)

    prs.save(output_path)
    print(f"🎯 PowerPoint generado en: {output_path}")
